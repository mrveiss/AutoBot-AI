# Copyright 2025-2026 mrveiss
# SPDX-License-Identifier: Apache-2.0
# AutoBot - AI-Powered Automation Platform
# Author: mrveiss
"""Heading/section structure on the canonical ExtractedDocument (#18826, owner option (a))."""

import dataclasses
import io

import pytest

from media.document.extraction import ExtractedDocument, extract_document, extract_docx
from media.document.sections import Block, BlockKind, ListItem, Section, SourceSpan, render_sections


def test_existing_producers_stay_compatible_sections_default_to_empty():
    assert ExtractedDocument(format="pdf", text="x").sections == ()
    assert extract_document(b"plain words").sections == ()


def test_render_sections_is_plain_text_with_tables_through_the_shared_renderer():
    sections = (
        Section(level=1, heading="Title", blocks=(Block(kind=BlockKind.PARAGRAPH, text="Intro."),)),
        Section(
            level=2,
            heading="Data",
            blocks=(
                Block(kind=BlockKind.TABLE, rows=(("Label", "Count"), ("bug", "578"))),
                Block(
                    kind=BlockKind.LIST,
                    items=(
                        ListItem("first", 0, True),
                        ListItem("nested", 1),
                        ListItem("second", 0, True),
                    ),
                ),
                Block(kind=BlockKind.CODE, text="  indented()\n"),
            ),
        ),
    )
    assert render_sections(sections) == (
        "Title\n\nIntro.\n\nData\n\nLabel | Count\nbug | 578\n\n1. first\n  - nested\n2. second\n\n  indented()\n"
    )


def test_extract_docx_text_is_unchanged_and_carries_no_sections_yet():
    """DOCX structure is #17985 (outline levels, not localised style names); text stays byte-identical."""
    pytest.importorskip("docx", reason="python-docx needed to synthesize a DOCX fixture")
    from docx import Document

    doc = Document()
    doc.add_heading("Chapter", level=1)
    doc.add_paragraph("Body one.")
    buffer = io.BytesIO()
    doc.save(buffer)

    extracted = extract_docx(buffer.getvalue())
    assert extracted.text == "Chapter\nBody one."
    assert extracted.sections == ()


def test_new_fields_default_so_existing_constructions_keep_working():
    section = Section(1, "Title", (Block(BlockKind.TABLE, rows=(("a", "b"),)),))
    block = section.blocks[0]
    assert section.heading_path == () and section.span is None
    assert block.header == () and block.code_language is None and block.span is None
    assert SourceSpan() == SourceSpan(page_start=None, page_end=None, time_start_s=None, time_end_s=None)


def test_span_round_trips_on_section_and_block():
    pages = SourceSpan(page_start=3, page_end=4)
    timed = SourceSpan(time_start_s=12.5, time_end_s=19.0)
    section = Section(2, "H", (Block(BlockKind.PARAGRAPH, text="x", span=timed),), heading_path=("T", "H"), span=pages)
    rebuilt = Section(**{f.name: getattr(section, f.name) for f in dataclasses.fields(Section)})
    assert rebuilt == section
    assert rebuilt.span.page_start == 3 and rebuilt.span.page_end == 4 and rebuilt.span.time_start_s is None
    assert rebuilt.blocks[0].span == SourceSpan(None, None, 12.5, 19.0)
    assert dataclasses.asdict(section)["span"] == {
        "page_start": 3,
        "page_end": 4,
        "time_start_s": None,
        "time_end_s": None,
    }


def test_render_sections_renders_a_table_header_first_then_its_body_rows():
    table = Block(kind=BlockKind.TABLE, header=("Label", "Count"), rows=(("bug", "578"), ("cleanup", "9")))
    sections = (Section(level=1, heading="Data", blocks=(table,), heading_path=("Data",)),)
    assert render_sections(sections) == "Data\n\nLabel | Count\nbug | 578\ncleanup | 9"


def test_header_and_body_render_like_the_same_rows_without_a_header():
    split = Block(kind=BlockKind.TABLE, header=("h1", "h2"), rows=(("a", "b"),))
    flat = Block(kind=BlockKind.TABLE, rows=(("h1", "h2"), ("a", "b")))
    assert render_sections((Section(0, "", (split,)),)) == render_sections((Section(0, "", (flat,)),))


def test_code_language_is_metadata_and_never_rendered():
    code = Block(kind=BlockKind.CODE, text="print(1)", code_language="python")
    assert render_sections((Section(0, "", (code,)),)) == "print(1)"
